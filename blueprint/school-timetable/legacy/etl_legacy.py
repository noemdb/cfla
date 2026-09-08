#!/usr/bin/env python3
"""
ETL: normaliza los Excel legacy de horarios/carga (año escolar 2025-2026)
a CSVs tabulares accesibles en blueprint/school-timetable/legacy/csv/.

Fuentes:
  HORARIOS MEDIA GENERAL.xlsx   (hojas 1ER..5TO AÑO, PROFESORES)
  HORARIOS PRIMARIA.xlsx         (hojas GRADOS — grados 1-6 e Inicial 1-3, ESPECIALISTAS)
  CARGA DISGREGADA ABRIL.xlsx    (fuente principal de carga, 23-04-2026)
  CARGA DISGREGADA INGLES.xlsx   (referencia área inglés, 24-03-2026)

Los PPTX de redes sociales son render visual del mismo horario y CARGA
ACADEMICA ABRIL.xlsx es la misma distribución en matriz; ambos fuera del ETL.

Estructura verificada de las grillas (MEDIA y PRIMARIA):
  - Encabezado 'HORAS' en columna A (fila h); título de la grilla en D..F{h-1}
    (izquierda) y K..M{h-1} (derecha).
  - Franjas horarias en filas siguientes (columna A), L-V en B..F e I..M.
  - RECESO etiquetado en las celdas de slot (merge horizontal).
  - Las CLASES ocupan celdas combinadas VERTICALES de 1-2 franjas: un bloque
    real = franja inicial + continuaciones (media: 80min = 2×40; primaria
    80/80/70/35min). El ETL resuelve los merges y reporta el bloque completo.
  - Horas PM en notación 12h ('01:05' = 13:05) → normalizadas a 24h.

Salidas (utf-8, en legacy/csv/):
  legacy_estructura_horaria.csv  Bloques de clase y recesos por nivel/turno
  legacy_horario_secciones.csv    Slots por sección (bloque completo + grupo paralelo)
  legacy_horario_docentes.csv     Slots por docente (clase/cargo + grupo)
  legacy_carga_docentes.csv       Carga disgregada por docente (abril + inglés)
"""

import csv
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

LEGACY = Path("/home/nuser/code/cfla/blueprint/school-timetable/legacy")
OUT = LEGACY / "csv"
OUT.mkdir(parents=True, exist_ok=True)

LEFT_COLS = ["B", "C", "D", "E", "F"]    # L-V grilla izquierda
RIGHT_COLS = ["I", "J", "K", "L", "M"]   # L-V grilla derecha
TITLE_LEFT = ["D", "E", "F"]             # columnas del título (grilla izquierda)
TITLE_RIGHT = ["K", "L", "M"]            # columnas del título (grilla derecha)
DAY_NAMES = {"LUNES", "MARTES", "MIERCOLES", "JUEVES", "VIERNES"}
MEDIA_YEARS = ["1ER AÑO", "2DO AÑO", "3ER AÑO", "4TO AÑO", "5TO AÑO"]

warn = []


def norm(s):
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().upper()


def clean(s):
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip()


def to24(hh, mm):
    return f"{int(hh):02d}:{mm}"


def minutes(h):
    hh, mm = h.split(":")
    return int(hh) * 60 + int(mm)


def parse_range(raw):
    """'07:00 A 07:40' | '10:45: A 11:20' | '12:25 A 01:05' → ('12:25','13:05').

    Notación 12h: horas 01:00-05:59 son PM (+12h). El colegio abre 07:00,
    así que cualquier hora < 06:00 es tarde (13:xx-17:xx).
    """
    if not raw:
        return None, None
    txt = clean(raw).upper().replace(" A ", "-")
    m = re.search(r"(\d{1,2}):?(\d{2}):?\s*[-]\s*(\d{1,2}):?(\d{2})", txt)
    if not m:
        return None, None
    h1h, m1, h2h, m2 = int(m.group(1)), m.group(2), int(m.group(3)), m.group(4)
    if h1h < 6:
        h1h += 12
    if h2h < 6:
        h2h += 12
    h1, h2 = to24(h1h, m1), to24(h2h, m2)
    if minutes(h2) <= minutes(h1):  # seguridad ante rangos invertidos
        return h1, None
    return h1, h2


# ─────────────────────────────────────────────────────────────
# Lectura de grillas
# ─────────────────────────────────────────────────────────────

class Grid:
    def __init__(self, sheet, header_row, title_left, title_right, franjas, merges):
        self.sheet = sheet
        self.header_row = header_row
        self.title_left = title_left
        self.title_right = title_right
        self.franjas = franjas            # [(fila, h1, h2)]
        self.franjas_by_row = {f[0]: f for f in franjas}
        self.merges = merges              # {(col_idx, row): (min_row, max_row)}

    def slot(self, col, fila):
        """Valor del slot si (col,fila) es celda cabecera de un bloque combinado
        (o celda simple). None si es continuación de un merge o está vacía."""
        c_idx = column_index_from_string(col)
        key = (c_idx, fila)
        if key in self.merges:
            mn, mx = self.merges[key]
            if mn != fila:
                return None, None  # continuación
            span = mx - mn  # franjas adicionales cubiertas
        else:
            span = 0
        v = self.sheet[f"{col}{fila}"].value
        if v is None or clean(v) == "" or norm(v) == "RECESO":
            return None, None
        return v, span

    def block_times(self, fila, span):
        h1 = self.franjas_by_row[fila][1]
        last = self.franjas_by_row[fila + span][2]
        return h1, last


def find_grids(sheet, max_row):
    return [r for r in range(1, max_row + 1) if norm(sheet[f"A{r}"].value) == "HORAS"]


def grid_title(sheet, header_row, cols, where):
    for c in cols:
        v = clean(sheet[f"{c}{header_row-1}"].value)
        nv = norm(v)
        if not v:
            continue
        if "HORARIO" in nv or nv in DAY_NAMES:
            continue
        return v
    return ""


def read_franjas(sheet, header_row, max_row):
    out = []
    r = header_row + 1
    blanks = 0
    while r <= max_row and blanks < 3:
        raw = sheet[f"A{r}"].value
        if raw is None or clean(raw) == "":
            blanks += 1
            r += 1
            continue
        if norm(raw) == "HORAS":
            break
        blanks = 0
        h1, h2 = parse_range(raw)
        if h1 and h2:
            out.append((r, h1, h2))
        r += 1
    return out


def merged_map(sheet):
    """{(col_idx, row): (min_row, max_row)} para merges VERTICALES de 1 columna."""
    m = {}
    for rng in sheet.merged_cells.ranges:
        if rng.min_col == rng.max_col:
            for row in range(rng.min_row, rng.max_row + 1):
                m[(rng.min_col, row)] = (rng.min_row, rng.max_row)
    return m


def load_grid(sheet, header_row, max_row, where):
    franjas = read_franjas(sheet, header_row, max_row)
    # Un lado sin título = no es una sección/docente: su contenido se
    # descarta (ruido de hoja, p. ej. celda suelta junto a INICIAL 3°).
    return Grid(sheet, header_row,
                grid_title(sheet, header_row, TITLE_LEFT, where),
                grid_title(sheet, header_row, TITLE_RIGHT, where),
                franjas, merged_map(sheet))


def iter_slots(grid, cols, title, nivel):
    """Yield (dia, h1, h2, valor) de cada bloque de clase de un lado de la grilla."""
    if not title:
        if grid.franjas:
            warn.append(f"grilla sin título ignorada ({nivel}, fila {grid.header_row})")
        return
    for fila, _, _ in grid.franjas:
        for day, col in enumerate(cols, start=1):
            v, span = grid.slot(col, fila)
            if v is None:
                continue
            if fila + span not in grid.franjas_by_row:
                span = 0  # merge que cae fuera de las franjas: trátalo como simple
            h1 = grid.franjas_by_row[fila][1]
            h2 = grid.franjas_by_row[fila + span][2]
            yield day, h1, h2, v


# ─────────────────────────────────────────────────────────────
# Paralelos (sub-grupos) y cargos
# ─────────────────────────────────────────────────────────────

GRUPO_PAREN = re.compile(r"^(.*?)\s*\((\d)\)\s*$", re.S)
GRUPO_TAG = re.compile(r"\(GRUPO\s*(\d)\)", re.I)
SEXO_TAG = re.compile(r"\((HEMBRAS|VARONES)\)\s*$", re.I)


def split_paralelo(raw):
    """Divide una celda con clases en paralelo (sub-grupos).

    'INGLÉS (GRUPO 1) FINANZAS (GRUPO 2)' → [('INGLÉS','1'), ('FINANZAS','2')]
    'INGLÉS (1)\\nROBÓTICA (2)'            → [('INGLÉS','1'), ('ROBÓTICA','2')]
    'EDUCACIÓN FÍSICA (HEMBRAS)'          → [('EDUCACIÓN FÍSICA','HEMBRAS')]
    'MATEMÁTICA'                          → [('MATEMÁTICA', None)]
    """
    if raw is None:
        return []
    t = clean(raw)
    if not t or norm(t) == "RECESO":
        return []
    tagged = GRUPO_TAG.findall(t)
    if len(tagged) >= 2:
        parts = re.split(r"\s*\(GRUPO\s*\d\)\s*", t, flags=re.I)
        materias = [re.sub(r"\s*\($", "", clean(p)).strip() for p in parts if clean(p)]
        return [(m, tagged[i] if i < len(tagged) else None)
                for i, m in enumerate(materias) if m]
    lines = [clean(x) for x in re.split(r"\n", str(raw)) if clean(x)]
    if len(lines) >= 2 and all(GRUPO_PAREN.match(x) for x in lines[:2]):
        out = []
        for x in lines:
            m = GRUPO_PAREN.match(x)
            if m and clean(m.group(1)):
                out.append((clean(m.group(1)), m.group(2)))
        return out
    m = SEXO_TAG.search(t)
    if m:
        base = clean(t[: m.start()]).strip()
        return [(base, m.group(1))] if base else []
    return [(t, None)]


CARGOS = ["PLANIFICACION", "ADMINISTRATIVAS", "APOYO", "GUIATURA", "O.V.", "OV"]

RECURSO_TAG = re.compile(r"\b(LABORATORIO|SALA|AUDITORIO|AULA)\b", re.I)


def grid_kind(title):
    """'docente' | 'recurso' — las grillas de PROFESORES incluyen aulas/labs."""
    return "recurso" if RECURSO_TAG.search(title or "") else "docente"


def clasificar_docente(raw):
    """Celda del horario de un docente → (contenido, grupo, tipo).

    '2DO B (1)'  → ('2DO B', '1', 'clase')
    'PLANIFICACIÓN' → ('PLANIFICACIÓN', None, 'cargo')
    """
    t = clean(raw)
    n = norm(t)
    tipo = "clase"
    for k in CARGOS:
        if k in n:
            tipo = "cargo"
            break
    grupo = None
    m = GRUPO_PAREN.match(t)
    if m:
        t, grupo = clean(m.group(1)), m.group(2)
    m = SEXO_TAG.search(t)
    if m:
        t, grupo = clean(t[: m.start()]).strip(), m.group(1)
    return t, grupo, tipo


# ─────────────────────────────────────────────────────────────
# 1+2. Estructura (bloques) y horario por sección
# ─────────────────────────────────────────────────────────────

def extract_media_secciones():
    rows = []
    blocks = set()
    wb = load_workbook(LEGACY / "HORARIOS MEDIA GENERAL.xlsx", data_only=True)
    for year in MEDIA_YEARS:
        sh = wb[year]
        for hr in find_grids(sh, 40):
            g = load_grid(sh, hr, 40, f"MEDIA {year}")
            for day, h1, h2, v in iter_slots(g, LEFT_COLS, g.title_left, "MEDIA"):
                blocks.add((h1, h2, False))
                for materia, grupo in split_paralelo(v):
                    rows.append(["MEDIA GENERAL", year, g.title_left, day, h1, h2, materia, grupo])
            for day, h1, h2, v in iter_slots(g, RIGHT_COLS, g.title_right, "MEDIA"):
                blocks.add((h1, h2, False))
                for materia, grupo in split_paralelo(v):
                    rows.append(["MEDIA GENERAL", year, g.title_right, day, h1, h2, materia, grupo])
    return rows, blocks


def extract_primaria_secciones():
    rows = []
    blocks = set()
    wb = load_workbook(LEGACY / "HORARIOS PRIMARIA.xlsx", data_only=True)
    sh = wb["GRADOS"]
    for hr in find_grids(sh, sh.max_row):
        g = load_grid(sh, hr, sh.max_row, "PRIMARIA GRADOS")
        for day, h1, h2, v in iter_slots(g, LEFT_COLS, g.title_left, "PRIMARIA"):
            blocks.add((h1, h2, False))
            for materia, grupo in split_paralelo(v):
                rows.append(["PRIMARIA", g.title_left, g.title_left, day, h1, h2, materia, grupo])
        for day, h1, h2, v in iter_slots(g, RIGHT_COLS, g.title_right, "PRIMARIA"):
            blocks.add((h1, h2, False))
            for materia, grupo in split_paralelo(v):
                rows.append(["PRIMARIA", g.title_right, g.title_right, day, h1, h2, materia, grupo])
    return rows, blocks


def collect_recesos(paths_sheets):
    """Franjas marcadas RECESO (merge horizontal de la etiqueta) → (h1, h2, True)."""
    out = {}
    for path, sheet_name in paths_sheets:
        wb = load_workbook(path, data_only=True)
        sh = wb[sheet_name]
        for hr in find_grids(sh, sh.max_row):
            franjas = read_franjas(sh, hr, sh.max_row)
            for fila, h1, h2 in franjas:
                if any(norm(sh[f"{c}{fila}"].value) == "RECESO" for c in LEFT_COLS + RIGHT_COLS):
                    out[(h1, h2)] = True
    return {(h1, h2, True) for (h1, h2) in out}


def write_estructura(blocks_media, blocks_primaria, recesos_media, recesos_primaria):
    rows = []
    for nivel, blocks, recesos in [
        ("MEDIA GENERAL", blocks_media, recesos_media),
        ("PRIMARIA", blocks_primaria, recesos_primaria),
    ]:
        for h1, h2, _ in sorted(blocks | recesos, key=lambda b: minutes(b[0])):
            is_break = (h1, h2, True) in recesos
            turno = "T" if minutes(h1) >= 13 * 60 else "M"
            rows.append([nivel, turno, h1, h2, minutes(h2) - minutes(h1),
                        "true" if is_break else "false"])
    with open(OUT / "legacy_estructura_horaria.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["nivel", "turno", "hora_inicio", "hora_fin", "minutos", "es_receso"])
        w.writerows(rows)
    return len(rows)


# ─────────────────────────────────────────────────────────────
# 3. Horario por docente
# ─────────────────────────────────────────────────────────────

def extract_docentes():
    rows = []
    wb = load_workbook(LEGACY / "HORARIOS MEDIA GENERAL.xlsx", data_only=True)
    sh = wb["PROFESORES"]
    for hr in find_grids(sh, sh.max_row):
        g = load_grid(sh, hr, sh.max_row, "MEDIA PROFESORES")
        for side_cols, title in [(LEFT_COLS, g.title_left), (RIGHT_COLS, g.title_right)]:
            if not title:
                continue
            for day, h1, h2, v in iter_slots(g, side_cols, title, "MEDIA PROFESORES"):
                contenido, grupo, tipo = clasificar_docente(v)
                if contenido:
                    rows.append(["MEDIA GENERAL", title, grid_kind(title), day, h1, h2, contenido, grupo or "", tipo])
    wb2 = load_workbook(LEGACY / "HORARIOS PRIMARIA.xlsx", data_only=True)
    sh = wb2["ESPECIALISTAS"]
    for hr in find_grids(sh, sh.max_row):
        g = load_grid(sh, hr, sh.max_row, "PRIMARIA ESPECIALISTAS")
        for side_cols, title in [(LEFT_COLS, g.title_left), (RIGHT_COLS, g.title_right)]:
            if not title:
                continue
            for day, h1, h2, v in iter_slots(g, side_cols, title, "PRIMARIA ESPECIALISTAS"):
                contenido, grupo, tipo = clasificar_docente(v)
                if contenido:
                    rows.append(["PRIMARIA", title, grid_kind(title), day, h1, h2, contenido, grupo or "", tipo])
    with open(OUT / "legacy_horario_docentes.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["nivel", "sujeto", "tipo_sujeto", "dia", "hora_inicio", "hora_fin", "contenido", "grupo", "tipo"])
        w.writerows(rows)
    return len(rows)


# ─────────────────────────────────────────────────────────────
# 4. Carga disgregada por docente
# ─────────────────────────────────────────────────────────────

def extract_carga():
    rows = []
    for fname, fuente in [
        ("CARGA DISGREGADA ABRIL.xlsx", "abril"),
        ("CARGA DISGREGADA INGLES.xlsx", "ingles"),
    ]:
        wb = load_workbook(LEGACY / fname, data_only=True)
        sh = wb["AÑO, GRADO Y AREA"]
        docente = area = None
        for r in range(1, sh.max_row + 1):
            num = sh[f"A{r}"].value
            doc = clean(sh[f"B{r}"].value)
            ar = clean(sh[f"C{r}"].value)
            grado = clean(sh[f"D{r}"].value)
            horas = sh[f"E{r}"].value
            total = sh[f"F{r}"].value
            obs = clean(sh[f"G{r}"].value)
            if doc:
                docente, area = doc, ar or area
            if norm(grado).startswith("TOTAL"):
                break
            if num is not None and doc:
                rows.append([fuente, num, docente, area or "", grado, horas or "", total or "", obs])
            elif grado or horas is not None:
                rows.append([fuente, "", docente or "", area or "", grado, horas or "", "", obs])
    with open(OUT / "legacy_carga_docentes.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["fuente", "n", "docente", "area", "grado_seccion", "horas", "total_docente", "observaciones"])
        w.writerows(rows)
    return len(rows)


def main():
    media_rows, media_blocks = extract_media_secciones()
    prim_rows, prim_blocks = extract_primaria_secciones()
    all_sec = media_rows + prim_rows
    with open(OUT / "legacy_horario_secciones.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["nivel", "grado", "seccion", "dia", "hora_inicio", "hora_fin", "materia", "grupo_paralelo"])
        w.writerows(all_sec)

    rec_media = collect_recesos([(LEGACY / "HORARIOS MEDIA GENERAL.xlsx", y) for y in MEDIA_YEARS])
    rec_prim = collect_recesos([(LEGACY / "HORARIOS PRIMARIA.xlsx", "GRADOS")])
    n_estructura = write_estructura(media_blocks, prim_blocks, rec_media, rec_prim)

    n_docentes = extract_docentes()
    n_carga = extract_carga()

    print(f"estructura={n_estructura} bloques | secciones={len(all_sec)} slots "
          f"({len(media_rows)} media + {len(prim_rows)} primaria) | "
          f"docentes={n_docentes} slots | carga={n_carga} filas")
    for w_ in warn:
        print(f"WARN: {w_}")
    print(f"outdir={OUT}")


if __name__ == "__main__":
    main()
